import os
import base64
from io import BytesIO

import streamlit as st
from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import Workbook
from pypdf import PdfReader
from docx import Document


# ============================================================
# CONFIGURATION
# ============================================================

load_dotenv()

API_KEY = os.getenv("OPENAI_API_KEY")

if not API_KEY:
    st.error(
        "OPENAI_API_KEY is not configured. "
        "Please add it to your .env file."
    )
    st.stop()

client = OpenAI(api_key=API_KEY)


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="AI QA Assistant",
    page_icon="🤖",
    layout="wide"
)


# ============================================================
# QA SYSTEM INSTRUCTIONS
# ============================================================

QA_INSTRUCTIONS = """
You are an expert Senior QA Engineer, SDET and Test Automation Engineer.

Your responsibilities include:

1. Functional testing
2. API testing
3. UI testing
4. Integration testing
5. Regression testing
6. Negative testing
7. Edge case testing
8. Boundary value analysis
9. Security testing
10. Performance testing
11. Test automation
12. Robot Framework
13. Selenium
14. Python
15. REST API testing
16. Defect analysis
17. Requirement analysis
18. UI/UX validation
19. Requirement traceability

Important rules:

- Do not invent undocumented application behavior.
- Do not invent exact status codes when they are not provided.
- Do not invent field length limits when they are not provided.
- Do not invent performance targets when they are not provided.
- Clearly identify assumptions.
- Think like a senior QA engineer.
- Cover positive, negative, edge and boundary scenarios.
- Consider validation and error handling.
- Consider authentication and authorization where applicable.
- Consider security risks where applicable.
- Consider data validation.
- Consider usability and UI behavior where applicable.
- Prefer practical, executable test scenarios.
- Avoid duplicate test cases.
- Keep test cases clear and traceable.
"""


# ============================================================
# SESSION STATE
# ============================================================

if "messages" not in st.session_state:
    st.session_state.messages = []

if "last_test_response" not in st.session_state:
    st.session_state.last_test_response = None


# ============================================================
# BASIC AI REQUEST
# ============================================================

def ask_ai(prompt, instructions=QA_INSTRUCTIONS):

    response = client.responses.create(
        model="gpt-6-luna",
        instructions=instructions,
        input=prompt
    )

    return response.output_text


# ============================================================
# AI REQUEST WITH IMAGE
# ============================================================

def ask_ai_with_image(
    prompt,
    image_bytes,
    image_type
):

    encoded_image = base64.b64encode(
        image_bytes
    ).decode("utf-8")

    image_data_url = (
        f"data:{image_type};base64,{encoded_image}"
    )

    response = client.responses.create(
        model="gpt-6-luna",
        instructions=QA_INSTRUCTIONS,
        input=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "input_text",
                        "text": prompt
                    },
                    {
                        "type": "input_image",
                        "image_url": image_data_url,
                        "detail": "high"
                    }
                ]
            }
        ]
    )

    return response.output_text


# ============================================================
# REQUIREMENT FILE EXTRACTION
# ============================================================

def extract_requirement(uploaded_file):

    file_name = uploaded_file.name.lower()

    try:

        # ----------------------------------------------------
        # TXT
        # ----------------------------------------------------

        if file_name.endswith(".txt"):

            return uploaded_file.read().decode(
                "utf-8",
                errors="ignore"
            )

        # ----------------------------------------------------
        # PDF
        # ----------------------------------------------------

        elif file_name.endswith(".pdf"):

            reader = PdfReader(uploaded_file)

            text = ""

            for page in reader.pages:

                page_text = page.extract_text()

                if page_text:

                    text += page_text + "\n"

            return text

        # ----------------------------------------------------
        # DOCX
        # ----------------------------------------------------

        elif file_name.endswith(".docx"):

            document = Document(uploaded_file)

            text = "\n".join(
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            )

            return text

        else:

            return None

    except Exception as e:

        st.error(
            f"Error reading file: {e}"
        )

        return None


# ============================================================
# EXCEL GENERATOR
# ============================================================

def create_excel_from_response(ai_response):

    workbook = Workbook()

    # ========================================================
    # SHEET 1 - TEST CASES
    # ========================================================

    test_sheet = workbook.active

    test_sheet.title = "Test Cases"

    test_headers = [
        "Test Case ID",
        "Category",
        "Scenario",
        "Test Data / Steps",
        "Expected Result"
    ]

    test_sheet.append(test_headers)

    lines = ai_response.splitlines()

    # --------------------------------------------------------
    # Parse Test Case table
    # --------------------------------------------------------

    for line in lines:

        clean_line = line.strip()

        if not clean_line.startswith("|"):
            continue

        if "---" in clean_line:
            continue

        parts = [
            part.strip()
            for part in clean_line.strip("|").split("|")
        ]

        if len(parts) < 5:
            continue

        test_id = parts[0]
        category = parts[1]
        scenario = parts[2]
        test_data = parts[3]
        expected = parts[4]

        # Skip table header
        if test_id.lower() in [
            "id",
            "test case id",
            "tc id"
        ]:
            continue

        # Only add actual TC rows
        if not test_id.upper().startswith("TC-"):
            continue

        test_sheet.append([
            test_id,
            category,
            scenario,
            test_data,
            expected
        ])

    # ========================================================
    # FORMAT TEST CASE SHEET
    # ========================================================

    for cell in test_sheet[1]:

        cell.font = cell.font.copy(
            bold=True
        )

    test_sheet.column_dimensions["A"].width = 18
    test_sheet.column_dimensions["B"].width = 25
    test_sheet.column_dimensions["C"].width = 55
    test_sheet.column_dimensions["D"].width = 65
    test_sheet.column_dimensions["E"].width = 65

    for row in test_sheet.iter_rows():

        for cell in row:

            cell.alignment = cell.alignment.copy(
                wrap_text=True,
                vertical="top"
            )

    test_sheet.freeze_panes = "A2"

    if test_sheet.max_row > 1:

        test_sheet.auto_filter.ref = (
            test_sheet.dimensions
        )

    # ========================================================
    # SHEET 2 - TRACEABILITY MATRIX
    # ========================================================

    trace_sheet = workbook.create_sheet(
        "Traceability Matrix"
    )

    trace_headers = [
        "Requirement ID",
        "Acceptance Criterion",
        "Figma/UI Evidence",
        "Test Case IDs",
        "Coverage Status",
        "Gap / Observation"
    ]

    trace_sheet.append(trace_headers)

    traceability_started = False

    for line in lines:

        clean_line = line.strip()

        # Detect traceability heading
        if (
            "traceability matrix" in
            clean_line.lower()
        ):

            traceability_started = True

            continue

        if not traceability_started:
            continue

        # Stop at next major section
        if (
            clean_line.startswith("##")
            and
            "traceability" not in
            clean_line.lower()
        ):

            break

        if not clean_line.startswith("|"):
            continue

        if "---" in clean_line:
            continue

        parts = [
            part.strip()
            for part in clean_line.strip("|").split("|")
        ]

        if len(parts) < 6:
            continue

        requirement_id = parts[0]
        acceptance_criterion = parts[1]
        figma_evidence = parts[2]
        test_case_ids = parts[3]
        coverage_status = parts[4]
        gap_observation = parts[5]

        # Skip header
        if requirement_id.lower() in [
            "requirement id",
            "id",
            "ac id"
        ]:
            continue

        # Only accept AC rows
        if not (
            requirement_id.upper().startswith("AC-")
            or
            requirement_id.upper().startswith("REQ-")
        ):
            continue

        trace_sheet.append([
            requirement_id,
            acceptance_criterion,
            figma_evidence,
            test_case_ids,
            coverage_status,
            gap_observation
        ])

    # ========================================================
    # FORMAT TRACEABILITY SHEET
    # ========================================================

    for cell in trace_sheet[1]:

        cell.font = cell.font.copy(
            bold=True
        )

    trace_sheet.column_dimensions["A"].width = 20
    trace_sheet.column_dimensions["B"].width = 55
    trace_sheet.column_dimensions["C"].width = 55
    trace_sheet.column_dimensions["D"].width = 25
    trace_sheet.column_dimensions["E"].width = 22
    trace_sheet.column_dimensions["F"].width = 65

    for row in trace_sheet.iter_rows():

        for cell in row:

            cell.alignment = cell.alignment.copy(
                wrap_text=True,
                vertical="top"
            )

    trace_sheet.freeze_panes = "A2"

    if trace_sheet.max_row > 1:

        trace_sheet.auto_filter.ref = (
            trace_sheet.dimensions
        )

    # ========================================================
    # SHEET 3 - QA SUMMARY
    # ========================================================

    summary_sheet = workbook.create_sheet(
        "QA Summary"
    )

    summary_sheet.append([
        "Metric",
        "Result"
    ])

    summary_sheet.append([
        "Analysis",
        "PBI + Figma"
    ])

    summary_sheet.append([
        "Test Cases Generated",
        "=COUNTA('Test Cases'!A:A)-1"
    ])

    summary_sheet.append([
        "Requirements in Traceability",
        "=COUNTA('Traceability Matrix'!A:A)-1"
    ])

    summary_sheet.append([
        "Covered Requirements",
        '=COUNTIF(\'Traceability Matrix\'!E:E,"Covered")'
    ])

    summary_sheet.append([
        "Partial Requirements",
        '=COUNTIF(\'Traceability Matrix\'!E:E,"Partial")'
    ])

    summary_sheet.append([
        "Not Covered Requirements",
        '=COUNTIF(\'Traceability Matrix\'!E:E,"Not Covered")'
    ])

    summary_sheet.append([
        "PBI/Figma Gaps",
        '=COUNTIF(\'Traceability Matrix\'!E:E,"Gap")'
    ])

    summary_sheet.append([
        "Note",
        "Coverage is based on AI analysis of the supplied PBI and Figma screenshot."
    ])

    for cell in summary_sheet[1]:

        cell.font = cell.font.copy(
            bold=True
        )

    summary_sheet.column_dimensions["A"].width = 40

    summary_sheet.column_dimensions["B"].width = 75

    for row in summary_sheet.iter_rows():

        for cell in row:

            cell.alignment = cell.alignment.copy(
                wrap_text=True,
                vertical="top"
            )

    # ========================================================
    # SAVE EXCEL
    # ========================================================

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output


# ============================================================
# SIDEBAR
# ============================================================

st.sidebar.title("🧰 QA Tools")

tool = st.sidebar.selectbox(
    "Select QA Tool",
    [
        "💬 General QA Chat",
        "🧪 Test Case Generator",
        "🔌 API Test Generator",
        "🐞 Bug Report Generator",
        "🤖 Robot Framework Generator",
        "🔍 Test Case Reviewer",
        "📄 Requirement Test Generator",
        "🎨 PBI + Figma QA Analyzer"
    ]
)


# ============================================================
# CLEAR CHAT
# ============================================================

if st.sidebar.button(
    "🗑️ Clear Chat"
):

    st.session_state.messages = []

    st.session_state.last_test_response = None

    st.rerun()


# ============================================================
# MAIN HEADER
# ============================================================

st.title(
    "🤖 AI QA Assistant"
)

st.caption(
    "AI-powered QA assistant for test design, API testing, "
    "bug reporting, automation and requirement analysis."
)


# ============================================================
# TOOL 1 - GENERAL QA CHAT
# ============================================================

if tool == "💬 General QA Chat":

    st.subheader(
        "💬 General QA Chat"
    )

    st.write(
        "Ask anything related to software testing, QA, "
        "automation, API testing or Robot Framework."
    )

    # Display previous messages
    for message in st.session_state.messages:

        with st.chat_message(
            message["role"]
        ):

            st.markdown(
                message["content"]
            )

    prompt = st.chat_input(
        "Ask your QA question..."
    )

    if prompt:

        st.session_state.messages.append(
            {
                "role": "user",
                "content": prompt
            }
        )

        with st.chat_message("user"):

            st.markdown(prompt)

        with st.chat_message("assistant"):

            with st.spinner(
                "Thinking..."
            ):

                try:

                    answer = ask_ai(prompt)

                    st.markdown(answer)

                    st.session_state.messages.append(
                        {
                            "role": "assistant",
                            "content": answer
                        }
                    )

                except Exception as e:

                    st.error(
                        f"Error while calling AI: {e}"
                    )


# ============================================================
# TOOL 2 - TEST CASE GENERATOR
# ============================================================

elif tool == "🧪 Test Case Generator":

    st.subheader(
        "🧪 Test Case Generator"
    )

    requirement = st.text_area(
        "Enter Requirement / Feature",
        height=220,
        placeholder=(
            "Example:\n"
            "User should be able to login using "
            "username and password."
        )
    )

    if st.button(
        "🚀 Generate Test Cases"
    ):

        if not requirement.strip():

            st.warning(
                "Please enter a requirement."
            )

        else:

            prompt = f"""
Analyze the following requirement.

Requirement:
{requirement}

Generate comprehensive test cases.

Cover:

1. Requirement Understanding
2. Assumptions
3. Positive Test Cases
4. Negative Test Cases
5. Edge Test Cases
6. Boundary Test Cases
7. Validation Test Cases
8. Security Test Cases
9. UI Test Cases if applicable
10. API Test Cases if applicable
11. Integration Test Cases if applicable

Use this exact table format:

| ID | Category | Scenario | Test Data / Steps | Expected Result |

Do not invent undocumented rules.
Clearly identify assumptions.
"""

            with st.spinner(
                "Generating test cases..."
            ):

                try:

                    answer = ask_ai(prompt)

                    st.session_state.last_test_response = answer

                    st.markdown(answer)

                except Exception as e:

                    st.error(
                        f"Error: {e}"
                    )


# ============================================================
# TOOL 3 - API TEST GENERATOR
# ============================================================

elif tool == "🔌 API Test Generator":

    st.subheader(
        "🔌 API Test Generator"
    )

    method = st.selectbox(
        "HTTP Method",
        [
            "GET",
            "POST",
            "PUT",
            "PATCH",
            "DELETE"
        ]
    )

    endpoint = st.text_input(
        "API Endpoint",
        placeholder="/api/login"
    )

    authentication = st.selectbox(
        "Authentication",
        [
            "None",
            "Bearer Token",
            "Basic Authentication",
            "API Key",
            "OAuth 2.0"
        ]
    )

    headers = st.text_area(
        "Request Headers",
        placeholder=(
            "Content-Type: application/json\n"
            "Accept: application/json"
        ),
        height=120
    )

    body = st.text_area(
        "Request Body",
        placeholder="""{
    "username": "test@test.com",
    "password": "Password123"
}""",
        height=180
    )

    if st.button(
        "🚀 Generate API Test Cases"
    ):

        if not endpoint.strip():

            st.warning(
                "Please enter the API endpoint."
            )

        else:

            prompt = f"""
Generate comprehensive API test cases.

API Details:

HTTP Method:
{method}

Endpoint:
{endpoint}

Authentication:
{authentication}

Request Headers:
{headers}

Request Body:
{body}

Generate:

1. Assumptions
2. Positive Test Cases
3. Negative Test Cases
4. Edge Test Cases
5. Boundary Test Cases
6. Authentication and Authorization Test Cases
7. Security Test Cases
8. Response Validation
9. Performance / Response Time Checks

Use this exact table format:

| ID | Category | Scenario | Test Data / Steps | Expected Result |

Important:

- Do not invent undocumented API behavior.
- Do not assume exact HTTP status codes unless documented.
- Do not invent field length limits.
- Do not invent performance targets.
- If something is unknown, state it as an assumption.
"""

            with st.spinner(
                "Generating API test cases..."
            ):

                try:

                    answer = ask_ai(prompt)

                    st.session_state.last_test_response = answer

                    st.markdown(answer)

                except Exception as e:

                    st.error(
                        f"Error: {e}"
                    )

    # --------------------------------------------------------
    # API EXCEL EXPORT
    # --------------------------------------------------------

    if st.session_state.last_test_response:

        st.divider()

        st.subheader(
            "📥 Export Test Cases"
        )

        if st.button(
            "📊 Prepare Excel File",
            key="api_prepare_excel"
        ):

            try:

                excel_file = create_excel_from_response(
                    st.session_state.last_test_response
                )

                st.download_button(
                    label="⬇️ Download Excel",
                    data=excel_file,
                    file_name="API_Test_Cases.xlsx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                    key="api_download_excel"
                )

            except Exception as e:

                st.error(
                    f"Excel creation failed: {e}"
                )


# ============================================================
# TOOL 4 - BUG REPORT GENERATOR
# ============================================================

elif tool == "🐞 Bug Report Generator":

    st.subheader(
        "🐞 Bug Report Generator"
    )

    description = st.text_area(
        "Describe the issue",
        height=200,
        placeholder=(
            "Example:\n"
            "When user selects a date from the date filter "
            "and clicks Apply, the selected date is not shown "
            "in the Logs table."
        )
    )

    if st.button(
        "🚀 Generate Bug Report"
    ):

        if not description.strip():

            st.warning(
                "Please describe the issue."
            )

        else:

            prompt = f"""
Convert the following issue description into a professional
software defect report.

Issue:
{description}

Generate:

1. Bug Title
2. Module
3. Severity
4. Priority
5. Environment
6. Preconditions
7. Steps to Reproduce
8. Actual Result
9. Expected Result
10. Impact
11. Regression Status
12. Suggested Evidence
13. Additional Notes

Do not invent facts.
If information is missing, mark it as "Not Provided".
"""

            with st.spinner(
                "Generating bug report..."
            ):

                try:

                    answer = ask_ai(prompt)

                    st.markdown(answer)

                except Exception as e:

                    st.error(
                        f"Error: {e}"
                    )


# ============================================================
# TOOL 5 - ROBOT FRAMEWORK GENERATOR
# ============================================================

elif tool == "🤖 Robot Framework Generator":

    st.subheader(
        "🤖 Robot Framework Generator"
    )

    test_requirement = st.text_area(
        "Describe the test you want to automate",
        height=220,
        placeholder=(
            "Example:\n"
            "Login to the application and verify "
            "that the dashboard is displayed."
        )
    )

    framework = st.selectbox(
        "Framework",
        [
            "Robot Framework + SeleniumLibrary",
            "Robot Framework + RequestsLibrary"
        ]
    )

    if st.button(
        "🚀 Generate Robot Framework Code"
    ):

        if not test_requirement.strip():

            st.warning(
                "Please enter the test requirement."
            )

        else:

            prompt = f"""
Generate Robot Framework automation code.

Requirement:
{test_requirement}

Framework:
{framework}

Follow these principles:

- Use clear test case names.
- Use reusable keywords.
- Use variables where appropriate.
- Keep locators maintainable.
- Add useful assertions.
- Add comments where useful.
- Follow Robot Framework syntax.
- Avoid hard-coded waits where possible.
- Prefer explicit waits.
- Keep code readable.

Return:

1. Test Case
2. Required Variables
3. Required Keywords
4. Full Robot Framework code
5. Explanation
"""

            with st.spinner(
                "Generating Robot Framework code..."
            ):

                try:

                    answer = ask_ai(prompt)

                    st.code(
                        answer,
                        language="robotframework"
                    )

                except Exception as e:

                    st.error(
                        f"Error: {e}"
                    )


# ============================================================
# TOOL 6 - TEST CASE REVIEWER
# ============================================================

elif tool == "🔍 Test Case Reviewer":

    st.subheader(
        "🔍 Test Case Reviewer"
    )

    test_cases = st.text_area(
        "Paste your test cases",
        height=300,
        placeholder=(
            "Paste your existing test cases here..."
        )
    )

    if st.button(
        "🔎 Review Test Cases"
    ):

        if not test_cases.strip():

            st.warning(
                "Please paste test cases."
            )

        else:

            prompt = f"""
Review the following test cases as a Senior QA Engineer.

Test Cases:
{test_cases}

Analyze:

1. Requirement coverage
2. Missing positive scenarios
3. Missing negative scenarios
4. Missing edge cases
5. Missing boundary cases
6. Missing security scenarios
7. Missing validation scenarios
8. Duplicate test cases
9. Ambiguous test cases
10. Incorrect expected results
11. Test data issues
12. Automation suitability

Provide:

- Overall assessment
- Issues found
- Missing scenarios
- Recommended improvements
- Improved test cases where necessary
"""

            with st.spinner(
                "Reviewing test cases..."
            ):

                try:

                    answer = ask_ai(prompt)

                    st.markdown(answer)

                except Exception as e:

                    st.error(
                        f"Error: {e}"
                    )


# ============================================================
# TOOL 7 - REQUIREMENT TEST GENERATOR
# ============================================================

elif tool == "📄 Requirement Test Generator":

    st.subheader(
        "📄 Requirement → Test Cases → Excel"
    )

    st.write(
        "Upload a requirement document and generate "
        "structured QA test cases."
    )

    uploaded_file = st.file_uploader(
        "📄 Upload Requirement",
        type=[
            "pdf",
            "docx",
            "txt"
        ],
        key="requirement_upload"
    )

    if uploaded_file:

        st.success(
            f"Uploaded: {uploaded_file.name}"
        )

        requirement_text = extract_requirement(
            uploaded_file
        )

        if requirement_text:

            st.subheader(
                "📖 Extracted Requirement"
            )

            with st.expander(
                "View Requirement Text",
                expanded=False
            ):

                st.text_area(
                    "Requirement",
                    requirement_text,
                    height=300,
                    key="extracted_requirement"
                )

            if st.button(
                "🚀 Generate Test Cases",
                key="requirement_generate"
            ):

                requirement_prompt = f"""
You are a Senior QA Engineer.

Analyze the following software requirement.

REQUIREMENT
===========

{requirement_text}

Generate:

1. Requirement Understanding
2. Assumptions
3. Ambiguities / Missing Information
4. Positive Test Cases
5. Negative Test Cases
6. Edge Test Cases
7. Boundary Test Cases
8. Validation Test Cases
9. Security Test Cases
10. UI Test Cases where applicable
11. API Test Cases where applicable
12. Integration Test Cases where applicable
13. Error Handling Test Cases where applicable

Use EXACTLY this table format:

| ID | Category | Scenario | Test Data / Steps | Expected Result |

Rules:

- Do not invent undocumented requirements.
- Do not invent exact field limits.
- Do not invent exact HTTP status codes.
- Do not invent unsupported business rules.
- Clearly mention assumptions.
- Do not duplicate scenarios.
- Make test cases practical and executable.
"""

                with st.spinner(
                    "Analyzing requirement..."
                ):

                    try:

                        answer = ask_ai(
                            requirement_prompt
                        )

                        st.session_state.last_test_response = answer

                        st.subheader(
                            "🧪 Generated Test Cases"
                        )

                        st.markdown(answer)

                    except Exception as e:

                        st.error(
                            f"Error: {e}"
                        )

        else:

            st.error(
                "Unable to extract requirement text."
            )

    if st.session_state.last_test_response:

        st.divider()

        st.subheader(
            "📥 Export Generated Test Cases"
        )

        if st.button(
            "📊 Prepare Excel File",
            key="requirement_prepare_excel"
        ):

            try:

                excel_file = create_excel_from_response(
                    st.session_state.last_test_response
                )

                st.download_button(
                    label="⬇️ Download Test Cases Excel",
                    data=excel_file,
                    file_name="Requirement_Test_Cases.xlsx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                    key="requirement_download_excel"
                )

            except Exception as e:

                st.error(
                    f"Excel creation failed: {e}"
                )


# ============================================================
# TOOL 8 - PBI + FIGMA QA ANALYZER
# ============================================================

elif tool == "🎨 PBI + Figma QA Analyzer":

    st.subheader(
        "🎨 PBI + Figma QA Analyzer"
    )

    st.info(
        "Use only dummy or sanitized requirements and UI "
        "screenshots. This tool does not connect to your "
        "company's Azure DevOps or Figma."
    )

    # ========================================================
    # PBI
    # ========================================================

    st.markdown(
        "### 📋 PBI / User Story"
    )

    pbi_text = st.text_area(
        "User Story / PBI Description",
        height=180,
        placeholder=(
            "Example:\n"
            "As an administrator, I want to add a device "
            "so that I can manage the device from the system."
        ),
        key="pbi_text"
    )

    # ========================================================
    # ACCEPTANCE CRITERIA
    # ========================================================

    st.markdown(
        "### ✅ Acceptance Criteria"
    )

    acceptance_criteria = st.text_area(
        "Acceptance Criteria",
        height=220,
        placeholder=(
            "Example:\n"
            "1. Device Name is mandatory.\n"
            "2. IP Address is mandatory.\n"
            "3. Save button creates the device.\n"
            "4. Cancel button closes the form.\n"
            "5. Invalid IP should show an error."
        ),
        key="acceptance_criteria"
    )

    # ========================================================
    # BUSINESS RULES
    # ========================================================

    st.markdown(
        "### 📌 Additional Business Rules"
    )

    business_rules = st.text_area(
        "Optional Business Rules",
        height=150,
        placeholder=(
            "Example:\n"
            "Device name must be unique.\n"
            "IP address must be valid."
        ),
        key="business_rules"
    )

    # ========================================================
    # FIGMA SCREENSHOT
    # ========================================================

    st.markdown(
        "### 🎨 Figma Design Screenshot"
    )

    figma_image = st.file_uploader(
        "Upload a screenshot/export from Figma",
        type=[
            "png",
            "jpg",
            "jpeg"
        ],
        key="figma_screenshot"
    )

    if figma_image:

        st.success(
            f"Uploaded: {figma_image.name}"
        )

        image_bytes = figma_image.getvalue()

        st.image(
            image_bytes,
            caption="Uploaded Figma Design",
            use_container_width=True
        )

    # ========================================================
    # GENERATE ANALYSIS
    # ========================================================

    if st.button(
        "🚀 Analyze PBI + Figma & Generate Test Cases",
        type="primary",
        key="pbi_figma_generate"
    ):

        if not pbi_text.strip():

            st.warning(
                "Please enter the PBI / User Story."
            )

        elif not acceptance_criteria.strip():

            st.warning(
                "Please enter the Acceptance Criteria."
            )

        elif not figma_image:

            st.warning(
                "Please upload a Figma screenshot."
            )

        else:

            prompt = f"""
You are a Senior QA Engineer performing a combined
requirements analysis and UI design review.

Analyze:

1. PBI / User Story
2. Acceptance Criteria
3. Business Rules
4. Uploaded Figma UI screenshot

==================================================
PBI / USER STORY
==================================================

{pbi_text}

==================================================
ACCEPTANCE CRITERIA
==================================================

{acceptance_criteria}

==================================================
BUSINESS RULES
==================================================

{business_rules}

==================================================
REQUIREMENT ANALYSIS
==================================================

Provide:

1. Requirement Understanding
2. UI Understanding
3. Assumptions
4. Missing Information
5. Potential PBI/Figma Mismatches
6. Potential UX/Usability Issues
7. Requirement Coverage Assessment

==================================================
FIGMA ANALYSIS
==================================================

Identify visible elements such as:

- Text fields
- Dropdowns
- Buttons
- Checkboxes
- Radio buttons
- Tables
- Tabs
- Filters
- Links
- Icons
- Labels
- Mandatory indicators
- Error messages
- Navigation elements
- Sections
- Search fields

Do not assume behavior that cannot be determined
from the screenshot.

==================================================
TEST CASE GENERATION
==================================================

Generate test cases covering:

1. Requirement
2. Positive
3. Negative
4. Edge
5. Boundary
6. Validation
7. UI
8. Security
9. Accessibility
10. Navigation
11. Error Handling
12. PBI-Figma Gap

Use EXACTLY this format:

| ID | Category | Scenario | Test Data / Steps | Expected Result |

Test Case IDs must be sequential:

TC-001
TC-002
TC-003
...

==================================================
REQUIREMENT TRACEABILITY MATRIX
==================================================

After the test cases, create a Requirement Traceability Matrix.

Assign IDs to acceptance criteria:

AC-01
AC-02
AC-03
...

Map every acceptance criterion to:

1. Figma/UI evidence
2. Test Case IDs
3. Coverage status
4. Gap / observation

Use EXACTLY:

| Requirement ID | Acceptance Criterion | Figma/UI Evidence | Test Case IDs | Coverage Status | Gap / Observation |

Coverage Status must be exactly one of:

Covered
Partial
Not Covered
Gap

Definitions:

Covered:
Requirement is represented in the UI/design and has
corresponding test cases.

Partial:
Requirement is partially represented or tested.

Not Covered:
Requirement exists but no suitable test case covers it.

Gap:
There is a potential mismatch or missing information
between PBI and Figma.

==================================================
QA COVERAGE SUMMARY
==================================================

After the traceability matrix provide:

- Total Acceptance Criteria
- Covered Requirements
- Partially Covered Requirements
- Not Covered Requirements
- PBI/Figma Gaps
- Overall QA Coverage Assessment

Do not invent numerical percentages.

If you provide a percentage, calculate it from
the traceability matrix.

==================================================
IMPORTANT RULES
==================================================

- Do not invent undocumented functionality.
- Do not invent exact field limits.
- Do not invent exact validation messages unless visible.
- Do not invent API behavior.
- Clearly identify assumptions.
- If PBI requires something not represented in Figma,
  identify it as a potential gap.
- If Figma shows something not mentioned in the PBI,
  identify it as a potential gap.
- Every acceptance criterion should have at least
  one corresponding test case where possible.
- Every test case should be traceable to a requirement
  or clearly identified as a supporting UI test.
- Avoid duplicate test cases.
"""

            with st.spinner(
                "AI is analyzing the PBI and Figma design..."
            ):

                try:

                    image_bytes = figma_image.getvalue()

                    file_name = (
                        figma_image.name.lower()
                    )

                    if file_name.endswith(".png"):

                        image_type = "image/png"

                    elif (
                        file_name.endswith(".jpg")
                        or
                        file_name.endswith(".jpeg")
                    ):

                        image_type = "image/jpeg"

                    else:

                        image_type = "image/png"

                    answer = ask_ai_with_image(
                        prompt,
                        image_bytes,
                        image_type
                    )

                    st.session_state.last_test_response = answer

                    st.divider()

                    st.subheader(
                        "🤖 AI QA Analysis"
                    )

                    st.markdown(answer)

                except Exception as e:

                    st.error(
                        f"Error while analyzing PBI + Figma: {e}"
                    )

    # ========================================================
    # PBI + FIGMA EXCEL EXPORT
    # ========================================================

    if st.session_state.last_test_response:

        st.divider()

        st.subheader(
            "📊 Export PBI + Figma QA Results"
        )

        if st.button(
            "📥 Prepare Excel File",
            key="pbi_figma_prepare_excel"
        ):

            try:

                excel_file = create_excel_from_response(
                    st.session_state.last_test_response
                )

                st.download_button(
                    label="⬇️ Download PBI + Figma Excel",
                    data=excel_file,
                    file_name="PBI_Figma_QA_Analysis.xlsx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                    key="pbi_figma_download_excel"
                )

            except Exception as e:

                st.error(
                    f"Excel creation failed: {e}"
                )


# ============================================================
# FOOTER
# ============================================================

st.sidebar.divider()

st.sidebar.caption(
    "🤖 AI QA Assistant"
)

st.sidebar.caption(
    "Python • Streamlit • OpenAI • QA Automation"
)